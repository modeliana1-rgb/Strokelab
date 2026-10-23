"""Presentación de la defensa del TFM StrokeLab sobre la plantilla UE 2022."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from PIL import Image

FIG = '/home/user/strokelab/memoria/figuras/'
ROJO, NEGRO, GRIS, GRIS_CL, BLANCO = (RGBColor(0xFF, 0x32, 0x28), RGBColor(0x1E, 0x1E, 0x1E), RGBColor(0x50, 0x4B, 0x4B),
                                      RGBColor(0xEE, 0xED, 0xED), RGBColor(0xFF, 0xFF, 0xFF))
AZUL, VERDE, ROSA = RGBColor(0x2D, 0x41, 0xB4), RGBColor(0x23, 0xB4, 0xAA), RGBColor(0xFF, 0xAF, 0xAA)

POST = []
p = Presentation('plantilla_UE_2022.pptx')
lay = {l.name: l for l in p.slide_layouts}

# Conservar solo la portada 1 (foto del campus) y borrar el resto de diapositivas de ejemplo
ids = p.slides._sldIdLst
for sid in list(ids)[1:]:
    p.part.drop_rel(sid.rId)
    ids.remove(sid)


def ph(s, idx):
    for x in s.placeholders:
        if x.placeholder_format.idx == idx:
            return x
    raise KeyError(idx)


def limpiar(s):
    """Quita los marcadores vacíos (si no, PowerPoint muestra «Haga clic para…»)."""
    for x in list(s.placeholders):
        if not x.has_text_frame or not x.text_frame.text.strip():
            x._element.getparent().remove(x._element)


def nueva(nombre, titulo=None, subt=None, notas=None, ancho_titulo=13.5):
    s = p.slides.add_slide(lay[nombre])
    if titulo is not None:
        t = s.shapes.title
        t.text = titulo
        t.width = Inches(min(ancho_titulo, 11.8))
        for q in t.text_frame.paragraphs:
            for r in q.runs:
                r.font.size = Pt(32)
    if subt is not None:
        ph(s, 18).text = subt
        ph(s, 18).width = Inches(13.5)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    return s


def texto(s, x, y, w, h, parrafos, tam=16, color=NEGRO, negrita=False, alin=PP_ALIGN.LEFT, ancla=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancla
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    for i, par in enumerate(parrafos):
        q = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        q.alignment = alin
        runs = par if isinstance(par, list) else [(par, {})]
        for txt, op in runs:
            r = q.add_run()
            r.text = txt
            r.font.size = Pt(op.get('tam', tam))
            r.font.bold = op.get('b', negrita)
            r.font.color.rgb = op.get('color', color)
        q.space_after = Pt(op.get('esp', 6) if runs else 6)
    return tb


def caja(s, x, y, w, h, relleno, linea=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = relleno
    if linea is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = linea
        r.line.width = Pt(1)
    r.shadow.inherit = False
    return r


def tarjeta(s, x, y, w, h, cabeza, cuerpo, relleno=GRIS_CL, ccab=ROJO, ccue=NEGRO, num=None):
    caja(s, x, y, w, h, relleno)
    y0 = y + 0.3
    if num:
        texto(s, x + 0.3, y0, w - 0.6, 0.6, num, tam=30, color=ccab, negrita=True)
        y0 += 0.65
    texto(s, x + 0.3, y0, w - 0.6, 0.8, cabeza, tam=18, color=ccab, negrita=True)
    texto(s, x + 0.3, y0 + 0.75, w - 0.6, h - (y0 - y) - 0.95, cuerpo, tam=14, color=ccue)


def cifra(s, x, y, w, numero, etiqueta, cnum=ROJO, cet=NEGRO, tam=40):
    texto(s, x, y, w, 0.85, numero, tam=tam, color=cnum, negrita=True)
    texto(s, x, y + 0.9, w, 0.9, etiqueta, tam=13, color=cet)


def imagen(s, ruta, x, y, w=None, h=None):
    W, H = Image.open(ruta).size
    if w and h:                                   # encajar manteniendo proporción y centrar
        k = min(w / W, h / H)
        ww, hh = W * k, H * k
        return s.shapes.add_picture(ruta, Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))
    return s.shapes.add_picture(ruta, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def pie(s, txt, y=7.95):
    texto(s, 0.99, y, 13.5, 0.35, txt, tam=11, color=GRIS)


def tabla(s, x, y, w, filas, anchos, tam=13, alto_fila=0.42, resaltar=()):
    t = s.shapes.add_table(len(filas), len(filas[0]), Inches(x), Inches(y), Inches(w), Inches(alto_fila * len(filas))).table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(filas):
        t.rows[i].height = Inches(alto_fila)
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(v)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = ROJO if i == 0 else (ROSA if i in resaltar else (GRIS_CL if i % 2 == 0 else BLANCO))
            for q in c.text_frame.paragraphs:
                for r in q.runs:
                    r.font.size = Pt(tam)
                    r.font.bold = i == 0 or i in resaltar
                    r.font.color.rgb = BLANCO if i == 0 else NEGRO
    return t


# ------------------------------------------------------------------ 1 · Portada
s = p.slides[0]
for x in s.placeholders:
    n = x.name
    if n.startswith('Título'):
        x.text = 'StrokeLab'
    elif n == 'Marcador de texto 7':
        x.text = 'Análisis biomecánico explicable de la técnica y la fatiga en natación'
    elif n == 'Marcador de texto 8':
        x.text = 'Diana Belem Cruz Montesinos'
    elif n == 'Marcador de texto 10':
        x.text = 'Dir.: Francisco Damian Maleno Gonzalez'
    elif n == 'Marcador de texto 9':
        x.text = 'Máster en IA · 24/10/2026'
for x in s.placeholders:
    if x.name in ('Marcador de texto 8', 'Marcador de texto 10', 'Marcador de texto 9'):
        x.left, x.width = Inches(3.99), Inches(3.33)
        x.top = Inches({'Marcador de texto 8': 5.55, 'Marcador de texto 10': 5.95, 'Marcador de texto 9': 6.3}[x.name])
        x.height = Inches(0.36)
        for q in x.text_frame.paragraphs:
            for r in q.runs:
                r.font.size = Pt(12 if x.name == 'Marcador de texto 8' else 10)
s.notes_slide.notes_text_frame.text = (
    'Buenos días. Soy Diana Cruz y presento StrokeLab, un sistema que analiza un vídeo de natación y le dice al '
    'entrenador cómo de eficiente es la brazada, en qué momento aparece la fatiga técnica y por qué. '
    'Todo funciona con una cámara normal y un portátil sin GPU. [≈ 20 s]')

# ------------------------------------------------------------------ 2 · Índice
s = nueva('Indice', 'Índice', notas='Esta es la estructura: el problema, el sistema, cómo lo validé, una demostración con '
          'dos grabaciones de 100 m libre y las conclusiones. [≈ 15 s]', ancho_titulo=6)
s.shapes.title.left, s.shapes.title.top = Inches(1.75), Inches(1.1)
s.shapes.title.height = Inches(1.2)
for r in s.shapes.title.text_frame.paragraphs[0].runs:
    r.font.size = Pt(44)
ph(s, 15).top = Inches(2.6)
ph(s, 15).left = Inches(1.75)
ph(s, 15).text = ''
tf = ph(s, 15).text_frame
for i, (n, t) in enumerate([('01', 'El problema y el objetivo'), ('02', 'StrokeLab: visión + datos tabulares'),
                            ('03', 'Validación y resultados'), ('04', 'Demostración: 100 m libre'),
                            ('05', 'Conclusiones y trabajo futuro')]):
    q = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = q.add_run(); r.text = f'{n}   '; r.font.color.rgb = ROJO; r.font.bold = True; r.font.size = Pt(24)
    r = q.add_run(); r.text = t; r.font.size = Pt(24)
    q.space_after = Pt(14)

# ------------------------------------------------------------------ 3 · Problema
s = nueva('Texto', '¿Cuándo aparece la fatiga y por qué?', 'Hoy el entrenador elige entre el ojo experto y los sensores caros',
          notas='Cuando un nadador se cansa, su técnica cambia antes de que caiga el tiempo: acorta la brazada, sube la '
          'frecuencia, baja el codo. Hoy esto se detecta a ojo, de forma subjetiva, o con sistemas instrumentados '
          '—sensores inerciales, varias cámaras— que son caros y complejos. StrokeLab propone una tercera vía: una sola '
          'cámara convencional y un portátil, y un sistema que no solo detecta la fatiga sino que explica qué variable '
          'la provoca. [≈ 1 min]')
tarjeta(s, 0.99, 3.2, 4.3, 3.3, 'Observación', 'Depende de la experiencia del entrenador. No cuantifica ni deja registro comparable entre sesiones.', num='01')
tarjeta(s, 5.6, 3.2, 4.3, 3.3, 'Sistemas instrumentados', 'Sensores inerciales o varias cámaras: coste elevado y montaje complejo en la piscina.', num='02')
tarjeta(s, 10.21, 3.2, 4.3, 3.3, 'StrokeLab', 'Una cámara convencional y un portátil sin GPU. Mide, detecta el inicio de la fatiga y explica su causa.',
        relleno=ROJO, ccab=BLANCO, ccue=BLANCO, num='03')
limpiar(s)

# ------------------------------------------------------------------ 4 · Objetivo
s = nueva('Texto', 'Objetivo: IA explicable desde un vídeo', 'Indicaciones del director y cómo se cumplen',
          notas='El objetivo general es un sistema de IA explicable que, desde vídeo monocular y en la CPU de un portátil, '
          'cuantifique la técnica, detecte el inicio de la fatiga y explique sus causas. El director marcó cuatro '
          'indicaciones: la explicabilidad es lo más importante, dos verticales de IA —visión y datos tabulares—, '
          'modos diferenciados y optimizar YOLO. La tabla resume cómo responde el sistema a cada una. [≈ 1 min]')
tabla(s, 0.99, 3.15, 14.0, [
    ['Indicación', 'Implementación en StrokeLab'],
    ['La explicabilidad (SHAP) es lo más importante', 'Isolation Forest sobre el estado fresco del propio nadador + TreeSHAP + explicación en texto'],
    ['Dos verticales de IA: visión y tabular', 'Visión: pose 2D (YOLOv8n) y 3D (MotionBERT) · Tabular: variables por ciclo y modelo de fatiga'],
    ['Modos diferenciados (vídeo, tabular, audio)', 'Vídeo y tabular implementados; el audio queda como trabajo futuro'],
    ['Optimizar YOLO (MoveNet, ViTPose)', 'Comparativa en CPU; YOLOv8n elegido; ViTPose como trabajo futuro'],
], [5.2, 8.8], tam=14, alto_fila=0.62)
limpiar(s)

# ------------------------------------------------------------------ 5 · Arquitectura
s = nueva('Texto', 'Dos verticales de IA',
          notas='Esta es la arquitectura. La vertical de visión estima la pose: 17 puntos del cuerpo por fotograma con '
          'YOLOv8n en 2D y MotionBERT en 3D, y los limpia con filtros anatómicos. La vertical tabular divide el nado en '
          'ciclos de brazada, calcula por ciclo la frecuencia, la distancia por ciclo y los ángulos de codo, hombro, '
          'cadera, rodilla y tronco, aprende el patrón fresco del propio nadador con Isolation Forest y explica cada '
          'desviación con SHAP. El resultado llega al entrenador como vídeo anotado, tablas y un texto. [≈ 1 min 15 s]')
imagen(s, FIG + 'arquitectura.png', 2.2, 2.25, 11.6, 6.0)
limpiar(s)

# ------------------------------------------------------------------ 6 · Visión
s = nueva('Texto', 'Adaptar la pose humana al agua', 'La pose bajo el agua necesita tres ajustes',
          notas='YOLO se entrenó con personas de pie, no con nadadores horizontales bajo el agua. Tres decisiones lo '
          'adaptan. Primero, el giro: pruebo el fotograma sin girar y girado 90 grados y me quedo con la orientación que '
          'mejor detecta. Segundo, filtros anatómicos: descarto cuerpos imposibles, como un tronco vertical en vista '
          'lateral o la cadera encima del hombro. Tercero, la elección del modelo: en el mismo vídeo YOLOv8n dio 15 '
          'ciclos válidos frente a 3 de MoveNet. El 3D de MotionBERT reconstruye mal las piernas de un nadador '
          'horizontal, así que en vista lateral uso los ángulos 2D. [≈ 1 min]')
tarjeta(s, 0.99, 3.2, 4.3, 3.3, 'Giro ±90°', 'Se prueba el fotograma girado y se conserva la orientación que mejor detecta; las coordenadas vuelven al original.', num='01')
tarjeta(s, 5.6, 3.2, 4.3, 3.3, 'Filtros anatómicos', 'Se descartan troncos verticales en vista lateral, segmentos imposibles y codos < 25°.', num='02')
tarjeta(s, 10.21, 3.2, 4.3, 3.3, 'YOLOv8n-Pose', '15 ciclos válidos frente a 3 de MoveNet en el mismo vídeo; 6,6 FPS en la CPU del portátil.', num='03')
pie(s, '3D (MotionBERT): válido en codo y cadera, no en rodilla → ángulos 2D en vista lateral y 3D en vista frontal u oblicua.', y=6.8)
limpiar(s)

# ------------------------------------------------------------------ 7 · Ciclos
s = nueva('Texto', 'Contar brazadas con la mano', 'La mano más profunda respecto al eje del cuerpo marca cada brazada',
          notas='Bajo el agua YOLO copia el brazo visible en el oculto: las dos muñecas salen casi idénticas. Por eso no '
          'sigo cada brazo por separado, sino la mano más profunda respecto al eje del cuerpo. Cada pico es una brazada '
          'y en crol dos brazadas forman un ciclo. Lo validé contra mi conteo manual: el error en la frecuencia está '
          'entre el 6,5 y el 8,2 %, en vista lateral y frontal. [≈ 1 min]')
imagen(s, FIG + 'aaron_brazadas.png', 0.99, 3.0, 8.4, 3.5)
tabla(s, 9.8, 3.1, 5.2, [['Secuencia (P1, crol)', 'Manual', 'Sistema', 'Error'],
                         ['GX011614 · lateral', '54,0', '50,5', '−6,5 %'],
                         ['IMG_7207 · frontal', '80,0', '73,9', '−7,6 %'],
                         ['IMG_7215 · frontal', '51,1', '55,2', '+8,2 %']], [2.1, 0.95, 1.1, 1.05], tam=11, alto_fila=0.5)
texto(s, 9.8, 5.3, 5.2, 0.6, 'Frecuencia de ciclo (ciclos/min) frente al conteo manual', tam=11, color=GRIS)
pie(s, 'Gris: intervalo del conteo manual (segundos 30-40). Triángulos: brazadas detectadas.', y=6.7)
limpiar(s)

# ------------------------------------------------------------------ 8 · Fatiga
s = nueva('Texto', 'Fatiga: desviación del patrón fresco', 'Aprendizaje no supervisado: no hacen falta etiquetas de fatiga',
          notas='No hay etiquetas de fatiga, así que planteo el problema como detección de anomalías respecto al propio '
          'nadador. El primer 30 % de sus ciclos es su patrón fresco. Isolation Forest puntúa cada ciclo por lo distinto '
          'que es de ese patrón. El inicio de la fatiga es el primer momento con tres ciclos seguidos por encima del '
          'percentil 95 de la fase fresca, y lo contrasto con el detector de cambios PELT. En la gráfica, con datos '
          'sintéticos y fatiga programada, el sistema la detecta al comienzo de la transición. [≈ 1 min 15 s]')
pasos = [('1', 'Fase fresca', 'Primer 30 % de ciclos del propio nadador'),
         ('2', 'Isolation Forest', 'Puntúa la anomalía de cada ciclo'),
         ('3', 'Inicio de la fatiga', '3 ciclos seguidos sobre el p95 fresco; contraste con PELT'),
         ('4', 'TreeSHAP', 'Qué variable empuja cada ciclo hacia la fatiga')]
for i, (n, c, d) in enumerate(pasos):
    y = 3.05 + i * 1.15
    caja(s, 0.99, y, 0.8, 0.8, ROJO)
    texto(s, 0.99, y, 0.8, 0.8, n, tam=24, color=BLANCO, negrita=True, alin=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE)
    texto(s, 2.0, y, 4.3, 0.4, c, tam=16, negrita=True)
    texto(s, 2.0, y + 0.42, 4.3, 0.6, d, tam=13, color=GRIS)
imagen(s, FIG + 'sintetico_fatiga_timeline.png', 6.7, 2.95, 8.3, 4.5)
pie(s, 'Gráfica: datos sintéticos con fatiga programada (referencia conocida), no un nadador real.', y=7.55)
limpiar(s)

# ------------------------------------------------------------------ 9 · SHAP
s = nueva('Texto', 'SHAP explica la fatiga', 'Atribución local en el ciclo de inicio (secuencia sintética de 90 s)',
          notas='La parte más importante para el director es la explicación. Para el ciclo en que empieza la fatiga, '
          'TreeSHAP reparte la puntuación de anomalía entre las variables. En la secuencia sintética de 90 s el inicio '
          'se detecta en el ciclo 35, a los 51,9 segundos, y las causas son el alcance del brazo derecho, que pasa de 1,78 a '
          '1,05 troncos, la asimetría de brazos, del 2 al 23 %, y la amplitud de patada. Son justo variables que se '
          'alteraron en la simulación: la explicación es verificable. PELT situó el cambio en el ciclo 33. [≈ 1 min]')
imagen(s, FIG + 'sintetico_shap_waterfall_inicio.png', 0.99, 2.9, 7.0, 5.0)
caja(s, 8.5, 3.05, 6.5, 2.4, NEGRO)
texto(s, 8.8, 3.3, 5.9, 2.0, [
    [('Lo que lee el entrenador', {'b': True, 'color': ROJO, 'tam': 16})],
    [('Inicio de la fatiga en el ciclo 35 (t = 51,9 s). Causas: alcance del brazo derecho (1,78 → 1,05 troncos), '
      'asimetría de brazos (2 % → 23 %) y amplitud de patada (0,31 → 0,21 troncos).', {'color': BLANCO, 'tam': 15})]])
texto(s, 8.5, 5.8, 6.5, 1.0, 'Las variables señaladas coinciden con las alteradas en la simulación: la explicación es verificable.',
      tam=14, color=NEGRO, negrita=True)
pie(s, 'Datos sintéticos de referencia conocida (memoria, Figura 7 y apartado 4.6.2).', y=7.95)
limpiar(s)

# ------------------------------------------------------------------ 10 · Resultados reales
s = nueva('Texto', 'Resultados con nadadores reales', '17 secuencias · 4 participantes · 4 estilos',
          notas='Con vídeo real analicé 17 secuencias de 4 participantes en 4 estilos: 6,6 minutos grabados, de los que '
          '2,2 minutos fueron analizables, con 95 ciclos válidos. La frecuencia se aproxima al conteo manual con un '
          'error del 6,5 al 8,2 %. La vista automática clasificó bien las 17 secuencias, aunque el umbral se fijó con ellas, y la familia del estilo '
          'en 12 de 13. Las sesiones de pasadas cortas no dieron fatiga, y tampoco falsas alarmas: es lo esperable, '
          'porque son pasadas a ritmo técnico con descansos. El nadador fue analizable en el 34 % del tiempo: la detección '
          'bajo el agua es la principal limitación. [≈ 1 min]')
items = [('95', 'ciclos válidos de 6,6 min grabados (2,2 min analizables)'),
         ('6,5-8,2 %', 'error de la frecuencia frente al conteo manual'),
         ('17/17', 'secuencias reales con la vista de cámara bien clasificada (umbral fijado con ellas)'),
         ('12/13', 'secuencias con la familia del estilo correcta (reglas biomecánicas)'),
         ('0', 'falsas alarmas de fatiga en 3 sesiones de pasadas cortas'),
         ('34 %', 'del tiempo grabado con el nadador analizable: la detección es el cuello de botella')]
for i, (n, e) in enumerate(items):
    x = 0.99 + (i % 3) * 4.7
    y = 3.05 + (i // 3) * 2.35
    caja(s, x, y, 4.4, 2.1, GRIS_CL if i != 0 else ROJO)
    cifra(s, x + 0.3, y + 0.2, 3.9, n, e, cnum=BLANCO if i == 0 else ROJO, cet=BLANCO if i == 0 else NEGRO, tam=36)
limpiar(s)

# ------------------------------------------------------------------ 11 · Sintéticos
s = nueva('Texto', 'Validación con referencia conocida', 'Datos sintéticos: se sabe exactamente cuándo empieza la fatiga',
          notas='Como no tengo una referencia real de fatiga, construí un generador de datos sintéticos en el que sé la '
          'verdad: la frecuencia, la velocidad y el momento exacto en que empieza la fatiga. Con él, el error mediano es '
          'del 4,7 % en frecuencia, del 1,3 % en velocidad y del 4,3 % en distancia por ciclo. El detector encuentra la '
          'fatiga en el 100 % de los casos con un 10 % de falsas alarmas, y en tiempo real detecta 10 de 10 sin falsas '
          'alarmas. Son datos sintéticos y así los presento: sirven para comprobar que el método funciona cuando la '
          'verdad se conoce. [≈ 1 min]')
for i, (n, e) in enumerate([('4,7 %', 'error mediano en la frecuencia de ciclo'), ('1,3 %', 'error mediano en la velocidad'),
                            ('4,3 %', 'error mediano en la distancia por ciclo')]):
    cifra(s, 0.99 + i * 2.6, 3.2, 2.4, n, e, tam=34)
caja(s, 0.99, 5.3, 7.4, 2.2, ROJO)
texto(s, 1.3, 5.5, 3.3, 1.9, [[('100 %', {'tam': 40, 'b': True, 'color': BLANCO})], [('detección de la fatiga · 10 % de falsas alarmas', {'tam': 13, 'color': BLANCO})]])
texto(s, 4.8, 5.5, 3.4, 1.9, [[('10/10', {'tam': 40, 'b': True, 'color': BLANCO})], [('en tiempo real, sin falsas alarmas', {'tam': 13, 'color': BLANCO})]])
imagen(s, FIG + 'caso_verdad_vs_sistema.png', 8.9, 2.9, 6.1, 4.7)
pie(s, 'Datos sintéticos (generador con referencia conocida), presentados como tales.', y=7.75)
limpiar(s)

# ------------------------------------------------------------------ 12 · Al entrenador
s = nueva('Texto', 'Del análisis al entrenador', 'Vídeo anotado y modo en tiempo real',
          notas='El resultado no es un número suelto. El vídeo anotado muestra en cada momento si el nadador está fresco o '
          'fatigado, el ciclo y la frecuencia, y una barra con la anomalía de cada ciclo y el umbral. El modo en tiempo '
          'real hace lo mismo mientras se reproduce el vídeo o la cámara, y muestra las causas de SHAP en directo. '
          'Además, el sistema decide solo la vista y la familia del estilo, y procesa listas de vídeos por lotes. [≈ 45 s]')
imagen(s, FIG + 'panel_video.png', 0.99, 2.95, 8.0, 4.5)
for i, (c, d) in enumerate([('Estado por ciclo', 'FRESCO / FATIGA, ciclo y frecuencia sobre el vídeo'),
                            ('Barra de anomalía', 'Puntuación de cada ciclo, umbral e inicio de la fatiga'),
                            ('Tiempo real', 'Calibra con los 10 primeros ciclos y explica con SHAP en directo'),
                            ('Automático', 'Vista de cámara, familia del estilo y lotes de vídeos')]):
    y = 3.0 + i * 1.12
    texto(s, 9.5, y, 5.5, 0.4, c, tam=16, color=ROJO, negrita=True)
    texto(s, 9.5, y + 0.42, 5.5, 0.6, d, tam=13)
pie(s, 'Captura del vídeo anotado sobre datos sintéticos.', y=7.55)
limpiar(s)

# ------------------------------------------------------------------ 12b · Demostración
s = nueva('Texto', 'Demostración: 100 m libre', 'Dos grabaciones nuevas analizadas con StrokeLab (GoPro, 60 fps)',
          notas='Para terminar, una demostración con dos grabaciones de 100 m libre, una suave y otra rápida, hechas después '
          'de entregar la memoria; no forman parte de sus resultados. [Reproducir el vídeo anotado del 100 suave.] '
          'En el suave el sistema decidió solo que la vista era lateral y el estilo crol, y obtuvo 20 ciclos válidos a 56 '
          'ciclos por minuto. En el rápido la cámara ve al nadador de frente u oblicuo gran parte del tiempo: en lateral '
          'solo salían 2 ciclos válidos y, indicando la vista frontal, 39 ciclos a 54 ciclos por minuto. En ninguno hay '
          'fatiga sostenida declarada. En el rápido, sin embargo, la puntuación de anomalía sube al final: los dos últimos ciclos son los más anómalos de todo el vídeo, pero la regla exige tres ciclos seguidos y el vídeo termina antes. La frecuencia de los dos es parecida; sin un conteo manual no '
          'afirmo cuál es más precisa. Es justo lo que dice la memoria: la vista y la detección condicionan el análisis. '
          '[≈ 1 min 30 s] NO mostrar los ángulos de cadera y rodilla del suave: las piernas se detectan mal en este vídeo.')
for k, (x, cab, its) in enumerate([
        (0.99, '100 suave', ['Vista lateral y estilo crol detectados automáticamente', '20 ciclos válidos · 42,5 s analizables de 86 s',
                             '56 ciclos/min de media', 'Sin fatiga sostenida']),
        (8.1, '100 rápido', ['Vista frontal u oblicua: indicada manualmente (en lateral, solo 2 ciclos)', '39 ciclos válidos · 76 s analizables de 106 s',
                             '54 ciclos/min de media', 'Sin fatiga declarada: los 2 últimos ciclos son los más anómalos del vídeo (la regla exige 3 seguidos)'])]):
    caja(s, x, 3.0, 6.9, 2.3, GRIS_CL, linea=RGBColor(0xB4, 0xB4, 0xB4))
    texto(s, x, 3.0, 6.9, 2.3, [[('Insertar aquí el vídeo anotado', {'tam': 13, 'color': GRIS})],
                                [('Insertar → Vídeo → Este dispositivo', {'tam': 11, 'color': GRIS})]], alin=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE)
    texto(s, x, 5.45, 6.9, 0.45, cab, tam=18, color=ROJO, negrita=True)
    texto(s, x, 5.95, 6.9, 1.5, [[('•  ' + t, {'tam': 13, 'esp': 4})] for t in its])
pie(s, 'Grabaciones posteriores a la entrega, solo como demostración; frecuencia sin validar con conteo manual.', y=7.6)
limpiar(s)

# ------------------------------------------------------------------ 13 · Posterior a la entrega
POST.append(len(p.slides._sldIdLst))
s = nueva('Texto', 'Tras la entrega: ajuste fino de la pose',
          '405 imágenes reales anotadas · validación cruzada por grabación',
          notas='Después de entregar quise atacar la principal limitación, la detección bajo el agua. Usé 405 imágenes '
          'reales de nadadores anotadas, públicas, y evalué con validación cruzada de cuatro pliegues separando por '
          'grabación para que no haya fuga de datos. El modelo ajustado detecta al nadador en el 78 % de las imágenes '
          'frente al 47 % del original, y el OKS sube de 0,10 a 0,40. Probé también modelos más grandes y datos '
          'sintéticos: no mejoran, lo que confirma que el límite son los datos reales del dominio. Y aplicado a mis '
          'vídeos aumenta el tiempo analizable un 26 %, pero la frecuencia no mejora: las reglas posteriores están '
          'calibradas con el modelo original y hay que recalibrarlas. Por eso el sistema de la memoria sigue con el '
          'modelo original. [≈ 1 min 30 s]')
cd = CategoryChartData()
cd.categories = ['Detección', 'PCK@0,1', 'PCK@0,2']
cd.add_series('Original (con giro)', (47, 9, 15))
cd.add_series('Ajustado', (78, 37, 56))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.99), Inches(3.2), Inches(7.6), Inches(4.4), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
ch.value_axis.maximum_scale = 100
ch.value_axis.has_major_gridlines = True
ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
ch.value_axis.tick_labels.font.size = Pt(11)
ch.category_axis.tick_labels.font.size = Pt(12)
for ser, col in zip(ch.plots[0].series, (RGBColor(0xB4, 0xB4, 0xB4), ROJO)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
pl = ch.plots[0]
pl.has_data_labels = True
pl.data_labels.number_format = '0"%"'
pl.data_labels.number_format_is_linked = False
pl.data_labels.font.size = Pt(12)
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.gap_width = 80
for i, (c, d) in enumerate([('OKS 0,10 → 0,40', 'Precisión de los puntos ×4 en imágenes nunca vistas'),
                            ('n ≈ s ≈ m', 'Modelos más grandes no mejoran: el límite son los datos'),
                            ('+ sintéticos = igual', 'SwimXYZ (777 imágenes) no mejora la prueba real'),
                            ('Vídeos propios', '+26 % tiempo analizable; la frecuencia no mejora → recalibrar')]):
    y = 3.0 + i * 1.15
    texto(s, 9.1, y, 5.9, 0.4, c, tam=16, color=ROJO, negrita=True)
    texto(s, 9.1, y + 0.42, 5.9, 0.6, d, tam=13)
pie(s, 'Roboflow «Deep Dive Pose Technique» (CC BY 4.0). Media de 4 pliegues. Resultados posteriores a la entrega de la memoria.', y=7.75)
limpiar(s)

# ------------------------------------------------------------------ 14 · Limitaciones y futuro
s = nueva('Texto', 'Limitaciones y trabajo futuro', 'La debilidad está en los datos, no en la arquitectura',
          notas='Las limitaciones son claras y las identifiqué con datos. Faltan grabaciones de nado continuo y una '
          'referencia fisiológica, así que la fatiga real no está validada todavía. El nadador solo es analizable el 34 % '
          'del tiempo y en vista lateral el modelo copia el brazo oculto. Con una cámara, los ángulos son proyecciones y '
          'el 3D falla en las piernas. El trabajo futuro sale directamente de ellas: series largas con lactato, esfuerzo '
          'percibido y valoración del entrenador; anotar fotogramas propios para medir OKS, PCK y mAP; ajustar la pose '
          'con SwimXYZ y datos propios; y ampliar con audio, sensores inerciales y una interfaz. [≈ 1 min]')
for k, (cab, its, x, rel) in enumerate([
        ('Limitaciones', ['Pocos datos reales y sin referencia fisiológica de la fatiga',
                          'Detección parcial: nadador analizable el 34 % del tiempo',
                          'Vista lateral: brazos y piernas de cada lado poco distinguibles',
                          'Una sola cámara: ángulos 2D proyectados; 3D poco fiable en piernas',
                          'Umbrales de vista y estilo fijados con los mismos vídeos'], 0.99, GRIS_CL),
        ('Trabajo futuro', ['Series largas de nado continuo con lactato, esfuerzo percibido y valoración experta',
                            'Anotar fotogramas propios: OKS, PCK y mAP con datos propios',
                            'Ajuste fino de la pose con SwimXYZ y fotogramas propios; ViTPose con GPU',
                            'Modo audio, sensores inerciales e interfaz para el entrenador'], 8.0, ROJO)]):
    caja(s, x, 3.05, 7.0, 4.5, rel)
    col = NEGRO if rel == GRIS_CL else BLANCO
    texto(s, x + 0.35, 3.3, 6.3, 0.5, cab, tam=20, color=ROJO if rel == GRIS_CL else BLANCO, negrita=True)
    texto(s, x + 0.35, 3.95, 6.3, 3.5, [[('•  ' + t, {'color': col, 'tam': 14, 'esp': 10})] for t in its])
limpiar(s)

# ------------------------------------------------------------------ 15 · Conclusiones
s = nueva('Texto', 'Conclusiones', '7 objetivos cumplidos y 2 parciales',
          notas='En resumen: siete objetivos cumplidos y dos parciales. Los parciales son honestos: la fatiga está validada '
          'con datos sintéticos y en tiempo real, pero falta la validación fisiológica con nadadores reales; y el '
          'clasificador de estilo supervisado no generalizaba, por eso lo sustituí por reglas biomecánicas. La lección '
          'principal del trabajo es que la calidad de los datos pesa más que el modelo: detecté una exactitud del 99,99 % '
          'que era fuga de datos y, desde entonces, contrasté cada resultado con una referencia. Gracias. [≈ 1 min]')
obj = [('OE1', 'Selección del modelo de pose', 'C'), ('OE2', 'Elevación a 3D', 'C*'), ('OE3', 'Ciclos de brazada', 'C'),
       ('OE4', 'Variables por ciclo', 'C'), ('OE5', 'Inicio de la fatiga', 'P'), ('OE6', 'Atribución con SHAP', 'C'),
       ('OE7', 'Clasificación del estilo', 'P'), ('OE8', 'Vídeo anotado y tiempo real', 'C'), ('OE9', 'Evaluación con datos sintéticos', 'C')]
for i, (c, d, g) in enumerate(obj):
    x = 0.99 + (i % 3) * 3.15
    y = 3.05 + (i // 3) * 1.45
    caja(s, x, y, 2.95, 1.25, ROJO if g.startswith('C') else NEGRO)
    texto(s, x + 0.2, y + 0.15, 2.6, 0.4, f'{c} · {"Cumplido" if g == "C" else ("Cumplido*" if g == "C*" else "Parcial")}', tam=12, color=BLANCO, negrita=True)
    texto(s, x + 0.2, y + 0.55, 2.6, 0.6, d, tam=13, color=BLANCO)
texto(s, 10.7, 3.05, 4.3, 4.2, [
    [('La calidad de los datos pesa más que el modelo.', {'tam': 20, 'b': True, 'color': ROJO, 'esp': 14})],
    [('Cada resultado se contrastó con una referencia: conteo manual o datos sintéticos con la verdad conocida.', {'tam': 14})]])
pie(s, '* Cumplido con limitaciones (3D válido en codo y cadera, no en rodilla).', y=7.55)
limpiar(s)

# ------------------------------------------------------------------ 16 · Contraportada
s = p.slides.add_slide(lay['Contraportada'])
s.shapes.title.text = 'Gracias'
ph(s, 16).text = '¿Preguntas?'
ph(s, 15).text = 'Diana Cruz'
ph(s, 22).text = 'StrokeLab · TFM Máster en IA'
ph(s, 23).text = 'Universidad Europea'
s.notes_slide.notes_text_frame.text = 'Muchas gracias. Quedo a disposición del tribunal para sus preguntas.'

# ================================================================== ANEXOS
s = p.slides.add_slide(lay['Blanco 01 rojo'])
s.shapes.title.text = 'Anexos'
ph(s, 18).text = 'Material de apoyo para las preguntas del tribunal'
limpiar(s)

s = nueva('Texto', 'A1 · Modelos de pose en CPU', 'Secuencia GX011614 (5120 × 2880 px), portátil de 4 hilos',
          notas='Si preguntan por qué YOLOv8n: es el más ligero sin perder detección, y en el vídeo real dio muchos más '
          'ciclos válidos que MoveNet. ViTPose se descartó por coste en CPU.')
tabla(s, 0.99, 3.15, 14.0, [['Modelo', 'Parámetros (M)', 'FPS en CPU', 'Detección (%)', 'Confianza'],
                            ['YOLOv8n-Pose (elegido)', '3,3', '6,6', '35,6', '0,76'],
                            ['YOLO11n-Pose', '2,9', '6,2', '37,6', '0,78'],
                            ['YOLOv8s-Pose', '11,6', '3,3', '35,6', '0,74'],
                            ['MoveNet Lightning', '—', 'No medido', '35-48 (GPU)', '0,34-0,37 (GPU)']],
      [4.6, 2.3, 2.3, 2.4, 2.4], tam=14, alto_fila=0.6, resaltar=(1,))
pie(s, 'La detección ronda el 36 % porque el nadador solo está en el encuadre parte del vídeo. En el mismo vídeo: 15 ciclos válidos con YOLO frente a 3 con MoveNet.', y=6.6)
limpiar(s)

POST.append(len(p.slides._sldIdLst))
s = nueva('Texto', 'B1 · Tamaño del modelo y sintéticos', 'Media de 4 pliegues por grabación · sin giro salvo indicación',
          notas='Tras el ajuste los tres tamaños quedan igual o el pequeño mejor: con unas 200 imágenes de entrenamiento '
          'el límite son los datos. Los sintéticos de SwimXYZ (vista frontal, fuera del agua) no mejoran la prueba, que '
          'es lateral bajo el agua, y aumentan las falsas detecciones.')
tabla(s, 0.99, 3.15, 14.0, [['Modelo', 'Detección', 'OKS', 'PCK@0,1', 'PCK@0,2', 'Falsos positivos (de 13)'],
                            ['YOLOv8n original (con giro, el del TFM)', '47 %', '0,10', '9 %', '15 %', '6'],
                            ['YOLOv8n ajustado', '78 %', '0,40', '37 %', '56 %', '3'],
                            ['YOLOv8s ajustado', '76 %', '0,37', '34 %', '54 %', '3,5'],
                            ['YOLOv8m ajustado', '69 %', '0,35', '33 %', '50 %', '2,5'],
                            ['YOLOv8n ajustado + 777 sintéticas (SwimXYZ)', '76 %', '0,39', '38 %', '54 %', '6,75']],
      [5.4, 1.6, 1.4, 1.6, 1.6, 2.4], tam=13, alto_fila=0.58, resaltar=(2,))
pie(s, 'Imágenes de Roboflow «Deep Dive Pose Technique» (CC BY 4.0), repartidas por grabación para evitar fuga de datos. Posterior a la entrega.', y=7.0)
limpiar(s)

POST.append(len(p.slides._sldIdLst))
s = nueva('Texto', 'B2 · Modelo ajustado en vídeos propios', '17 secuencias del TFM, misma resolución (640 px)',
          notas='Con el modelo ajustado el sistema analiza más tiempo y saca más ciclos, pero la frecuencia frente al '
          'conteo manual no mejora: la señal de la mano tiene picos dobles. Las reglas posteriores están calibradas con '
          'el modelo original; recalibrarlas exige más conteos manuales. Por eso el sistema del TFM mantiene el original.')
tabla(s, 0.99, 3.15, 7.0, [['', 'Original', 'Ajustado'], ['Tiempo analizable', '132 s', '166 s'], ['Ciclos válidos', '95', '109']],
      [2.8, 2.1, 2.1], tam=14, alto_fila=0.6)
tabla(s, 8.4, 3.15, 6.6, [['Frecuencia vs. manual', 'Original', 'Ajustado'], ['GX011614 (54,0)', '−6,5 %', '+35 %'],
                          ['IMG_7207 (80,0)', '−7,6 %', '−24 %'], ['IMG_7215 (51,1)', '+8,2 %', '+7 %']],
      [2.8, 1.9, 1.9], tam=14, alto_fila=0.6)
texto(s, 0.99, 5.9, 14.0, 1.4, 'Conclusión: un mejor modelo de pose no basta; hay que recalibrar el conteo de brazadas, la vista y el estilo con su nueva señal.',
      tam=16, negrita=True)
limpiar(s)

POST.append(len(p.slides._sldIdLst))
s = nueva('Texto', 'B3 · Vídeos públicos de televisión', 'Robustez fuera de las condiciones de diseño',
          notas='Probé dos vídeos de televisión. En la retransmisión del 800 libre de 2001, desde fuera del agua, con '
          'varios nadadores y cambios de plano, el sistema falla: la frecuencia salta de 40 a 90 ciclos por minuto. En '
          'un vídeo con cámara en el fondo y un solo nadador, detecta sola la vista frontal y sigue al nadador. El '
          'sistema está pensado para grabaciones de entrenamiento con un nadador.')
tarjeta(s, 0.99, 3.1, 6.8, 3.3, 'Retransmisión de carrera (fuera del agua)',
        'Varios nadadores y cambios de plano: el esqueleto salta de uno a otro y la frecuencia varía de 40 a 90 ciclos/min. No fiable.', num='✗')
tarjeta(s, 8.2, 3.1, 6.8, 3.3, 'Cámara en el fondo, un nadador',
        'Vista frontal detectada automáticamente y nadador seguido. 18 s: insuficiente para fatiga; frecuencia pendiente de conteo manual.',
        relleno=ROJO, ccab=BLANCO, ccue=BLANCO, num='✓')
limpiar(s)

s = nueva('Texto', 'A2 · Fatiga en las sesiones reales', 'Sesiones de pasadas cortas del mismo participante, estilo y cámara',
          notas='Ninguna sesión mostró fatiga ni falsas alarmas. Son pasadas cortas a ritmo técnico con descansos, no un '
          'esfuerzo continuo. No demuestra la capacidad de detectar fatiga real, pero sí que el sistema no inventa fatiga.')
tabla(s, 0.99, 3.15, 14.0, [['Sesión', 'Secuencias', 'Ciclos', 'Resultado'],
                            ['P1 · crol · GoPro', '3', '10', 'Sin fatiga sostenida'],
                            ['P1 · crol · móvil', '2', '7', 'Insuficiente (< 8 ciclos)'],
                            ['P4 · braza · GoPro', '2', '27', 'Sin fatiga sostenida'],
                            ['P4 · braza · móvil', '2', '4', 'Insuficiente (< 8 ciclos)'],
                            ['P3 · mariposa · GoPro', '2', '26', 'Sin fatiga sostenida']],
      [5.0, 2.4, 2.0, 4.6], tam=14, alto_fila=0.58)
limpiar(s)

s = nueva('Texto', 'A3 · Caso sintético completo (Tabla 7)', 'Secuencia de 120 s con fatiga programada entre t = 59 s y t = 81 s',
          notas='Si preguntan por la validación con referencia conocida: en vista lateral se estiman frecuencia, velocidad y '
          'distancia por ciclo con errores medianos del 4,7, 1,3 y 4,3 %; de frente, la frecuencia con un 0,3 %. El inicio '
          'se detecta a los 57,3 s, al comienzo de la transición, y SHAP señala variables alteradas en la simulación.')
tabla(s, 0.99, 3.1, 14.0, [['Medida', 'Vista lateral', 'Vista frontal'],
                           ['Ciclos analizados', '91', '98'],
                           ['Error mediano de la frecuencia de ciclo', '4,7 %', '0,3 %'],
                           ['Error mediano de la velocidad', '1,3 %', 'No observable de frente'],
                           ['Error mediano de la distancia por ciclo', '4,3 %', 'No observable de frente'],
                           ['Inicio de la fatiga detectado', 't = 57,3 s (ciclo 42)', 't = 57,8 s (ciclo 45)'],
                           ['Variables con mayor atribución SHAP', 'Apertura del hombro, asimetría de brazos, inclinación del tronco',
                            'Frecuencia de ciclo, asimetría de brazos, alcance del brazo']],
      [4.6, 4.7, 4.7], tam=13, alto_fila=0.58)
pie(s, 'Datos sintéticos. Fuente: memoria, Tabla 7.', y=7.4)
limpiar(s)

p.save('completa.pptx')
print('ok', len(p.slides), 'posteriores', POST)


def quedarse(entrada, salida, quitar):
    q = Presentation(entrada)
    ids = q.slides._sldIdLst
    for k, sid in reversed(list(enumerate(ids))):
        if k in quitar:
            q.part.drop_rel(sid.rId)
            ids.remove(sid)
    q.save(salida)


n = len(p.slides._sldIdLst)
quedarse('completa.pptx', 'StrokeLab_defensa.pptx', set(POST))
quedarse('completa.pptx', 'StrokeLab_posterior_entrega.pptx', set(range(n)) - set(POST) - {0})
